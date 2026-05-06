from __future__ import annotations

import re
import subprocess
from html.parser import HTMLParser
from urllib.parse import urlparse
from urllib.request import Request, urlopen
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from pypdf import PdfReader

from app.schemas import ParsedDocument, ParsedSection

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class _ReadableHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self._skip_depth = 0
        self.chunks: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "noscript", "svg"}:
            self._skip_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript", "svg"} and self._skip_depth:
            self._skip_depth -= 1

    def handle_data(self, data: str) -> None:
        if self._skip_depth:
            return
        normalized = re.sub(r"\s+", " ", data).strip()
        if normalized:
            self.chunks.append(normalized)


def detect_language(text: str) -> str:
    cjk = len(re.findall(r"[\u4e00-\u9fff]", text))
    ascii_letters = len(re.findall(r"[A-Za-z]", text))
    return "zh" if cjk >= ascii_letters else "en"


def split_sections(text: str) -> list[ParsedSection]:
    blocks = [line.strip() for line in text.splitlines()]
    sections: list[ParsedSection] = []
    current_heading = "导言"
    current_paragraphs: list[str] = []

    for block in blocks:
        if not block:
            continue
        if re.match(r"^(#+\s+.+|(?:\d+(?:\.\d+)*)[\s、.].+|Abstract|Introduction|Methods?|Results?|Discussion|Conclusion|References)$", block, re.I):
            if current_paragraphs:
                sections.append(ParsedSection(heading=current_heading, paragraphs=current_paragraphs))
            current_heading = re.sub(r"^#+\s*", "", block)
            current_paragraphs = []
        else:
            current_paragraphs.append(block)

    if current_paragraphs:
        sections.append(ParsedSection(heading=current_heading, paragraphs=current_paragraphs))

    if not sections:
        chunks = [chunk.strip() for chunk in re.split(r"\n{2,}", text) if chunk.strip()]
        return [ParsedSection(heading=f"片段 {index + 1}", paragraphs=[chunk]) for index, chunk in enumerate(chunks[:6])]
    return sections[:8]


def extract_citations(text: str) -> list[str]:
    citations = re.findall(r"\(([^()]{6,80}?\d{4}[^()]*)\)", text)
    numbered = re.findall(r"\[(\d{1,3}(?:,\s*\d{1,3})*)\]", text)
    return list(dict.fromkeys(citations + numbered))[:8]


def extract_formulas(text: str) -> list[str]:
    formulas = re.findall(r"([A-Za-z][A-Za-z0-9_]*\s*=\s*[^。\n]{3,60})", text)
    return list(dict.fromkeys(formulas))[:6]


def guess_doc_type(text: str, sections: list[ParsedSection]) -> tuple[str, dict[str, float]]:
    lowered = text.lower()
    paper_score = sum(keyword in lowered for keyword in ["abstract", "introduction", "method", "result", "discussion", "conclusion", "references"])
    paper_score += sum("摘要" in s.heading or "方法" in s.heading or "结论" in s.heading for s in sections)
    argument_score = sum(keyword in lowered for keyword in ["because", "therefore", "however", "claim", "evidence", "所以", "因此", "但是", "论证", "前提"])
    notes_score = sum(keyword in lowered for keyword in ["definition", "example", "exercise", "知识点", "定义", "例题", "章节"])
    notes_score += sum(len(section.paragraphs) <= 2 for section in sections)

    scores = {
        "paper": paper_score + 1.0,
        "argument_text": argument_score + 1.0,
        "notes_or_textbook": notes_score + 1.0,
    }
    total = float(sum(scores.values())) or 1.0
    normalized = {key: round(value / total, 3) for key, value in scores.items()}
    doc_type = max(normalized, key=normalized.get)
    return doc_type, normalized


def parse_plain_text(text: str, *, source_name: str = "文本输入", parse_strategy: str = "direct_text") -> ParsedDocument:
    normalized = text.replace("\r\n", "\n").strip()
    sections = split_sections(normalized)
    doc_type, doc_scores = guess_doc_type(normalized, sections)
    metadata = {
        "title": next((section.heading for section in sections if section.heading and section.heading != "导言"), source_name),
        "source_name": source_name,
        "character_count": len(normalized),
        "doc_type_scores": doc_scores,
    }
    return ParsedDocument(
        metadata=metadata,
        sections=sections,
        paragraphs=[paragraph for section in sections for paragraph in section.paragraphs][:30],
        citations=extract_citations(normalized),
        tables=[],
        formulas=extract_formulas(normalized),
        language=detect_language(normalized),
        doc_type_guess=doc_type,  # type: ignore[arg-type]
        parse_strategy=parse_strategy,
    )


def extract_readable_html(html: str) -> str:
    parser = _ReadableHTMLParser()
    parser.feed(html)
    return "\n".join(parser.chunks)


def fetch_web_text(url: str) -> str:
    parsed = urlparse(url.strip())
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise ValueError("请输入有效的 http 或 https 网页地址。")
    if parsed.hostname in {"localhost", "127.0.0.1", "::1"}:
        raise ValueError("暂不支持抓取本机地址。")

    request = Request(url, headers={"User-Agent": "PixelMaterialLearningConverter/0.1"})
    with urlopen(request, timeout=8) as response:  # noqa: S310 - user-provided learning URL, constrained by scheme above.
        content_type = response.headers.get("content-type", "")
        raw = response.read(1_200_000)
    charset_match = re.search(r"charset=([\w-]+)", content_type, re.I)
    encoding = charset_match.group(1) if charset_match else "utf-8"
    html = raw.decode(encoding, errors="ignore")
    if "html" in content_type.lower() or "<html" in html.lower():
        return extract_readable_html(html)
    return html


def parse_web_url(url: str) -> ParsedDocument:
    text = fetch_web_text(url)
    if len(text.strip()) < 120:
        raise ValueError("网页可提取文本过少，请换一个正文更完整的链接或粘贴文本。")
    return parse_plain_text(text, source_name=url.strip(), parse_strategy="web_url")


def parse_docx(path: Path) -> ParsedDocument:
    document = Document(path)
    paragraphs = [para.text.strip() for para in document.paragraphs if para.text.strip()]
    text = "\n".join(paragraphs)
    return parse_plain_text(text, source_name=path.name, parse_strategy="docx_text")


def _run_command(command: list[str]) -> str:
    result = subprocess.run(command, check=False, capture_output=True, text=True)
    return result.stdout.strip() if result.returncode == 0 else ""


def parse_pdf(path: Path) -> ParsedDocument:
    page_texts: list[str] = []
    try:
        reader = PdfReader(str(path))
        for page in reader.pages[:12]:
            page_texts.append((page.extract_text() or "").strip())
    except Exception:
        page_texts = []

    text = "\n".join(chunk for chunk in page_texts if chunk)
    strategy = "pdf_text_layer"

    if len(text) < 400:
        extracted = _run_command(["pdftotext", "-layout", str(path), "-"])
        if extracted:
            text = extracted
            strategy = "pdftotext"

    if len(text) < 250:
        with TemporaryDirectory() as temp_dir:
            prefix = Path(temp_dir) / "page"
            _run_command(["pdftoppm", "-png", "-f", "1", "-l", "5", str(path), str(prefix)])
            ocr_chunks: list[str] = []
            for image in sorted(Path(temp_dir).glob("page-*.png"))[:5]:
                ocr_text = _run_command(["tesseract", str(image), "stdout", "-l", "eng+chi_sim"])
                if ocr_text:
                    ocr_chunks.append(ocr_text)
            if ocr_chunks:
                text = "\n".join(ocr_chunks)
                strategy = "pdf_ocr"

    if len(text.strip()) < 120:
        raise ValueError("文档可提取文本过少，请上传更清晰的 PDF 或补充文本。")

    return parse_plain_text(text, source_name=path.name, parse_strategy=strategy)


def parse_file(path: Path, content_type: str | None) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".docx":
        return parse_docx(path)
    if suffix in {".txt", ".md"}:
        return parse_plain_text(path.read_text(encoding="utf-8", errors="ignore"), source_name=path.name, parse_strategy="plain_file")
    raise ValueError(f"暂不支持的文件类型：{suffix or content_type or 'unknown'}")
